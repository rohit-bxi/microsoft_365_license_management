# -*- coding: utf-8 -*-
import re

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError

GUID_RE = re.compile(r'^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$', re.I)
DOMAIN_RE = re.compile(r'^(?=.{4,253}$)([a-z0-9]([a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}$',
                       re.I)
TRUE_VALUES = ('1', 'true', 'yes', 'y')
FALSE_VALUES = ('0', 'false', 'no', 'n')
TEMPLATE_NAME = 'm365_tenant_import_template.xlsx'

# Importable m365.tenant fields: (field name, value type, required). The column header
# is the field label (or the technical name); blank optional cells keep the default.
COLUMNS = [
    ('name', 'text', True),
    ('domain', 'domain', True),
    ('tenant_id', 'guid', True),
    ('client_id', 'guid', True),
    ('client_secret', 'text', False),
    ('authority_url', 'url', False),
    ('graph_url', 'url', False),
    ('graph_scope', 'url', False),
    ('allow_license_operations', 'bool', False),
    ('auto_approve_admins', 'bool', False),
    ('sync_enabled', 'bool', False),
    ('sync_users', 'bool', False),
    ('company_id', 'company', False),
    ('low_availability_percent', 'percent', False),
    ('low_availability_units', 'units', False),
]
# Must be unique (case insensitive) within the file and among the existing tenants
UNIQUE_FIELDS = ('tenant_id', 'domain', 'name')
# Columns of the downloadable template; the other columns may be added to the file
TEMPLATE_FIELDS = ('name', 'domain', 'tenant_id', 'client_id', 'client_secret',
                   'authority_url', 'graph_url', 'graph_scope', 'allow_license_operations',
                   'auto_approve_admins', 'sync_enabled', 'sync_users')
# Imported tenants are fully enabled unless the file says otherwise
IMPORT_DEFAULTS = {
    'allow_license_operations': True,
    'auto_approve_admins': True,
    'sync_enabled': True,
    'sync_users': True,
}
# Example rows of the template (valid, importable test data with placeholder IDs)
SAMPLE_ROWS = [
    {'name': 'Contoso Ltd', 'domain': 'contoso.com',
     'tenant_id': 'aaaaaaaa-0000-0000-0000-000000000001',
     'client_id': 'bbbbbbbb-0000-0000-0000-000000000001',
     'client_secret': 'replace-with-client-secret',
     'allow_license_operations': True, 'auto_approve_admins': True,
     'sync_enabled': True, 'sync_users': True},
    # Read-only inventory, no user directory
    {'name': 'Fabrikam Inc', 'domain': 'fabrikam.com',
     'tenant_id': 'aaaaaaaa-0000-0000-0000-000000000002',
     'client_id': 'bbbbbbbb-0000-0000-0000-000000000002',
     'client_secret': 'replace-with-client-secret',
     'allow_license_operations': False, 'auto_approve_admins': False,
     'sync_enabled': True, 'sync_users': False},
    # Sovereign cloud (US Government) endpoints
    {'name': 'Northwind Traders', 'domain': 'northwindtraders.com',
     'tenant_id': 'aaaaaaaa-0000-0000-0000-000000000003',
     'client_id': 'bbbbbbbb-0000-0000-0000-000000000003',
     'client_secret': 'replace-with-client-secret',
     'authority_url': 'https://login.microsoftonline.us',
     'graph_url': 'https://graph.microsoft.us/v1.0',
     'graph_scope': 'https://graph.microsoft.us/.default',
     'allow_license_operations': True, 'auto_approve_admins': False,
     'sync_enabled': True, 'sync_users': True},
    # Required columns only: blank secret (added later), default URLs, switches TRUE
    {'name': 'Adventure Works', 'domain': 'adventure-works.com',
     'tenant_id': 'aaaaaaaa-0000-0000-0000-000000000004',
     'client_id': 'bbbbbbbb-0000-0000-0000-000000000004'},
]


class M365TenantImportWizard(models.TransientModel):
    _name = 'm365.tenant.import.wizard'
    _inherit = 'm365.file.import.mixin'
    _description = 'Import Microsoft 365 Tenants'

    file = fields.Binary(string="File", attachment=False,
                         help="CSV or Excel (.xlsx) file with one tenant per row.")
    filename = fields.Char()
    template_file = fields.Binary(attachment=False)
    line_ids = fields.One2many('m365.tenant.import.line', 'wizard_id', string="Rows")
    line_count = fields.Integer(compute='_compute_counts')
    error_count = fields.Integer(compute='_compute_counts')

    @api.depends('line_ids.status')
    def _compute_counts(self):
        for wiz in self:
            wiz.line_count = len(wiz.line_ids)
            wiz.error_count = len(wiz.line_ids.filtered(lambda l: l.status == 'error'))

    @api.onchange('file')
    def _onchange_file(self):
        # A new file must be validated again before tenants can be created
        self.line_ids = [Command.clear()]

    # ------------------------------------------------------------------
    # File reading
    # ------------------------------------------------------------------
    def _map_header(self, header):
        """Map column positions to m365.tenant field names."""
        Tenant = self.env['m365.tenant']
        lookup = {}
        for fname, _kind, _required in COLUMNS:
            lookup[fname] = fname
            lookup[self._normalise(Tenant._fields[fname].string)] = fname
        mapping, unknown = {}, []
        for index, title in enumerate(header):
            if not title:
                continue
            fname = lookup.get(self._normalise(title))
            if not fname:
                unknown.append(title)
            elif fname in mapping.values():
                raise UserError(_("The column \"%s\" appears more than once.", title))
            else:
                mapping[index] = fname
        if unknown:
            raise UserError(_(
                "Unknown column(s): %(unknown)s.\nAllowed columns: %(allowed)s.",
                unknown=", ".join(unknown),
                allowed=", ".join(Tenant._fields[f].string for f, _k, _r in COLUMNS)))
        missing = [Tenant._fields[f].string for f, _k, required in COLUMNS
                   if required and f not in mapping.values()]
        if missing:
            raise UserError(_("Missing required column(s): %s.", ", ".join(missing)))
        return mapping

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------
    def _convert(self, kind, text, companies):
        """Return the field value for a cell, raise ValueError when it is not valid."""
        if kind == 'guid':
            if not GUID_RE.match(text):
                raise ValueError(_("\"%s\" is not a valid ID (expected "
                                   "00000000-0000-0000-0000-000000000000).", text))
            return text.lower()
        if kind == 'domain':
            if not DOMAIN_RE.match(text):
                raise ValueError(_("\"%s\" is not a valid domain name.", text))
            return text.lower()
        if kind == 'url':
            if not text.lower().startswith('https://'):
                raise ValueError(_("\"%s\" must start with https://.", text))
            return text
        if kind == 'bool':
            if text.lower() in TRUE_VALUES:
                return True
            if text.lower() in FALSE_VALUES:
                return False
            raise ValueError(_("\"%s\" is not a Yes/No value.", text))
        if kind in ('percent', 'units'):
            try:
                number = float(text.replace(',', '.'))
            except ValueError:
                raise ValueError(_("\"%s\" is not a number.", text)) from None
            if kind == 'percent':
                if not 0 <= number <= 100:
                    raise ValueError(_("%s must be between 0 and 100.", text))
                return number
            if number < 0 or not number.is_integer():
                raise ValueError(_("%s must be a whole number of 0 or more.", text))
            return int(number)
        if kind == 'company':
            company = companies.get(text.lower())
            if not company:
                raise ValueError(_("Company \"%s\" not found in your allowed companies.", text))
            return company.id
        return text

    def _parse(self):
        """Read and validate the uploaded file.

        :return: list of dicts with the serial number (1 = first row below the
                 header), the raw cell texts, the m365.tenant values and the error
                 messages of each data row
        """
        rows = self._read_rows()
        if len(rows) < 2:
            raise UserError(_("The file contains no tenant rows below the header."))
        mapping = self._map_header(rows[0])
        Tenant = self.env['m365.tenant']
        kinds = {fname: kind for fname, kind, _required in COLUMNS}
        required = {fname for fname, _kind, req in COLUMNS if req}
        companies = {c.name.lower(): c for c in self.env.user.company_ids}
        # The unique constraint is global: also check archived / other company tenants
        tenants = Tenant.sudo().with_context(active_test=False).search([])
        existing = {fname: {(value or '').lower() for value in tenants.mapped(fname)}
                    for fname in UNIQUE_FIELDS}
        seen = {fname: {} for fname in UNIQUE_FIELDS}
        result = []
        for sn, values in enumerate(rows[1:], start=1):
            raw, vals, errors = {}, dict(IMPORT_DEFAULTS), []
            for index, fname in mapping.items():
                text = raw[fname] = values[index] if index < len(values) else ''
                label = Tenant._fields[fname].string
                if not text:
                    if fname in required:
                        errors.append(_("%s is required.", label))
                    continue
                try:
                    vals[fname] = self._convert(kinds[fname], text, companies)
                except ValueError as exc:
                    errors.append(_("%(field)s: %(error)s", field=label, error=exc))
            for fname in UNIQUE_FIELDS:
                key = (vals.get(fname) or '').lower()
                if not key:
                    continue
                label = Tenant._fields[fname].string
                if key in existing[fname]:
                    errors.append(_("%(field)s \"%(value)s\" already exists in Odoo.",
                                    field=label, value=vals[fname]))
                elif key in seen[fname]:
                    errors.append(_("Duplicate %(field)s, same as SN. %(sn)s.",
                                    field=label, sn=seen[fname][key]))
                seen[fname].setdefault(key, sn)
            result.append({'sn': sn, 'raw': raw, 'vals': vals, 'errors': errors})
        return result

    # ------------------------------------------------------------------
    # Buttons
    # ------------------------------------------------------------------
    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'name': _("Import Tenants"),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_validate(self):
        self.ensure_one()
        parsed = self._parse()
        self.line_ids = [Command.clear()] + [Command.create({
            'sn': row['sn'],
            'name': row['raw'].get('name'),
            'domain': row['raw'].get('domain'),
            'directory_id': row['raw'].get('tenant_id'),
            'client_id': row['raw'].get('client_id'),
            'status': 'error' if row['errors'] else 'valid',
            'message': "\n".join(row['errors']) or False,
        }) for row in parsed]
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_("Validate the file before creating the tenants."))
        # Parse again: the file or the existing tenants may have changed since validation
        parsed = self._parse()
        invalid = [row for row in parsed if row['errors']]
        if invalid:
            raise UserError(_(
                "No tenant has been created. Correct these rows and validate the file "
                "again:\n%s", "\n".join("SN. %s: %s" % (row['sn'], " ".join(row['errors']))
                                         for row in invalid)))
        if self.error_count:
            raise UserError(_("Validate the file again before creating the tenants."))
        tenants = self.env['m365.tenant'].create([row['vals'] for row in parsed])
        for tenant in tenants:
            tenant._message_log(body=_("Tenant created by bulk import from %s.",
                                       self.filename))
        # Do not keep client secrets in the wizard any longer than needed
        self.file = False
        action = self.env['ir.actions.act_window']._for_xml_id(
            'm365_license_management.action_m365_tenant')
        action.update(name=_("Imported Tenants"), domain=[('id', 'in', tenants.ids)],
                      target='main')
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _("Microsoft 365"),
                'message': _("%s tenant(s) created. Test their connection to start "
                             "synchronising.", len(tenants)),
                'type': 'success',
                'next': action,
            },
        }

    def _template_hints(self):
        return {
            'text': _("Text"),
            'domain': _("Domain name, e.g. contoso.com"),
            'guid': _("GUID, e.g. 11111111-2222-3333-4444-555555555555"),
            'url': _("https:// address"),
            'bool': _("TRUE or FALSE"),
        }

    def action_download_template(self):
        self.ensure_one()
        Tenant = self.env['m365.tenant']
        hints = self._template_hints()
        defaults = dict(Tenant.default_get(list(TEMPLATE_FIELDS)), **IMPORT_DEFAULTS)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = _("Tenants")
        required_fill = PatternFill('solid', fgColor='FFF2CC')
        bool_choice = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
        sheet.add_data_validation(bool_choice)
        columns = [column for column in COLUMNS if column[0] in TEMPLATE_FIELDS]
        for col, (fname, kind, required) in enumerate(columns, start=1):
            label = Tenant._fields[fname].string
            cell = sheet.cell(row=1, column=col, value=label)
            cell.font = Font(bold=True)
            hint = hints[kind]
            if required:
                cell.fill = required_fill
                hint = _("Required. %s", hint)
            elif fname in defaults:
                hint = _("%(hint)s. Blank = %(default)s", hint=hint, default=defaults[fname])
            else:
                hint = _("Optional. %s", hint)
            cell.comment = Comment(hint, 'Odoo')
            if kind == 'bool':
                bool_choice.add('%(c)s2:%(c)s1000' % {'c': cell.column_letter})
            sheet.column_dimensions[cell.column_letter].width = max(16, len(label) + 4)
        for sample in SAMPLE_ROWS:
            sheet.append([sample.get(fname) for fname, _kind, _required in columns])
        sheet.freeze_panes = 'A2'
        return self._download_workbook(book, TEMPLATE_NAME)


class M365TenantImportLine(models.TransientModel):
    _name = 'm365.tenant.import.line'
    _description = 'Microsoft 365 Tenant Import Row'
    _order = 'sn'

    wizard_id = fields.Many2one('m365.tenant.import.wizard', required=True,
                                ondelete='cascade')
    sn = fields.Integer(string="SN.")
    name = fields.Char(string="Tenant Name")
    domain = fields.Char(string="Primary Domain")
    directory_id = fields.Char(string="Directory (Tenant) ID")
    client_id = fields.Char(string="Application (Client) ID")
    status = fields.Selection([('valid', 'Valid'), ('error', 'Error')], string="Status")
    message = fields.Text(string="Errors")
