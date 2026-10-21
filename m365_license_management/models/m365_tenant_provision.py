# -*- coding: utf-8 -*-
"""Creation of new Microsoft 365 customer tenants through Microsoft Partner Center.

Odoo version of the ``Bulk-M365Customers.ps1`` provisioning script: every customer of
``Customers.csv`` becomes a ``m365.tenant.provision`` record that is checked and then
created with the Partner Center REST API (the API behind ``New-PartnerCustomer``). The
result of each customer (status, customer ID, error, timestamp) is kept on its record
instead of ``ProvisioningResults.csv``.

=========================================  ===========================================
Step                                       Partner Center endpoint (app-only auth)
=========================================  ===========================================
Test the connection                        GET  /v1/profiles/legalbusiness
Domain availability                        HEAD /v1/domains/{domain}
Address validation                         POST /v1/validations/address
Create the customer (tenant + admin user)  POST /v1/customers
=========================================  ===========================================

Reference: https://learn.microsoft.com/en-us/partner-center/developer/create-a-customer
"""
import re
import time
import uuid
from datetime import timedelta

import openpyxl
import requests
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from odoo import Command, _, api, fields, models
from odoo.exceptions import AccessError, UserError
from odoo.tools import config, email_normalize

from .m365_tenant import M365GraphError

ADMIN_GROUP = 'm365_license_management.group_m365_admin'
PC_APPLICATION = 'Odoo Microsoft 365 License Management'
INITIAL_DOMAIN = '.onmicrosoft.com'
DOMAIN_RE = re.compile(r'^[a-z0-9]([a-z0-9-]{0,61}[a-z0-9])?\.onmicrosoft\.com$')
CULTURE_RE = re.compile(r'^[a-z]{2}-[a-z]{2}$', re.I)
LANGUAGE_RE = re.compile(r'^[a-z]{2}$', re.I)
VERIFIED_ADDRESS = ('Verified', 'VerifiedShippable')
# Countries / regions where Microsoft requires the organization registration number
REGISTRATION_COUNTRIES = {
    'AM', 'AZ', 'BY', 'HU', 'KZ', 'KG', 'MD', 'RU', 'TJ', 'UZ', 'UA', 'BR', 'IN', 'ZA',
    'PL', 'AE', 'SA', 'TR', 'TH', 'VN', 'MM', 'IQ', 'SS', 'VE', 'CN',
}
REQUIRED_FIELDS = ('name', 'domain', 'first_name', 'last_name', 'email', 'address_line1',
                   'city', 'postal_code', 'country_id', 'culture', 'language')
# Data sent to Microsoft: a checked customer must be checked again when it changes
CUSTOMER_FIELDS = set(REQUIRED_FIELDS) | {
    'phone', 'address_line2', 'address_state', 'organization_registration_number',
    'enable_gdap', 'partner_account_id', 'address_confirmed',
}
PROVISION_BATCH = 20  # customers created per run of the scheduled action


class M365PartnerCenterTimeout(M365GraphError):
    """Partner Center did not answer: retry the call with the same MS-RequestId."""


class M365PartnerAccount(models.Model):
    """Microsoft Partner Center (CSP) account the customer tenants are created with."""
    _name = 'm365.partner.account'
    _description = 'Microsoft Partner Center Account'
    _inherit = ['mail.thread']
    _order = 'name'

    name = fields.Char(required=True, tracking=True)
    active = fields.Boolean(default=True)
    company_id = fields.Many2one('res.company', string="Company",
                                 default=lambda self: self.env.company)
    partner_tenant_id = fields.Char(
        string="Partner Directory (Tenant) ID", required=True, copy=False, tracking=True,
        help="Directory ID (or primary domain) of your Partner Center (CSP) tenant.")
    client_id = fields.Char(string="Application (Client) ID", required=True, copy=False,
                            tracking=True)
    client_secret = fields.Char(string="Client Secret", copy=False, groups=ADMIN_GROUP)
    authority_url = fields.Char(string="Authority URL", required=True,
                                default='https://login.microsoftonline.com')
    api_url = fields.Char(string="Partner Center API URL", required=True,
                          default='https://api.partnercenter.microsoft.com')
    api_scope = fields.Char(string="Partner Center Scope", required=True,
                            default='https://api.partnercenter.microsoft.com/.default')
    access_token = fields.Char(copy=False, groups='base.group_system')
    token_expiry = fields.Datetime(copy=False, groups='base.group_system')
    state = fields.Selection([
        ('draft', 'Not Tested'),
        ('connected', 'Connected'),
        ('error', 'Connection Error'),
    ], string="Status", default='draft', readonly=True, copy=False, tracking=True)
    partner_name = fields.Char(string="Partner Organisation", readonly=True, copy=False)
    last_error = fields.Text(string="Last Error", readonly=True, copy=False)
    default_culture = fields.Char(string="Default Culture", required=True, default='en-US',
                                  help="Culture of the new customers, e.g. en-US or en-IN.")
    default_language = fields.Char(string="Default Language", required=True, default='en',
                                   help="Two letter language code of the new customers.")
    enable_gdap = fields.Boolean(
        string="Create GDAP Relationship", default=True,
        help="Ask Microsoft to create the default granular delegated admin (GDAP) "
             "relationship with each new customer, so that you can administer it.")
    provision_ids = fields.One2many('m365.tenant.provision', 'partner_account_id',
                                    string="Customer Tenants")
    provision_count = fields.Integer(compute='_compute_provision_count')

    @api.depends('provision_ids')
    def _compute_provision_count(self):
        for rec in self:
            rec.provision_count = len(rec.provision_ids)

    def write(self, vals):
        # Changing the credentials invalidates the cached token
        if {'partner_tenant_id', 'client_id', 'client_secret', 'authority_url',
                'api_scope'} & set(vals):
            vals = dict(vals, access_token=False, token_expiry=False, state='draft')
        return super().write(vals)

    # ------------------------------------------------------------------
    # Partner Center client
    # ------------------------------------------------------------------
    @staticmethod
    def _format_error(resp):
        """Partner Center errors: ``{"code", "errorName", "description"}``."""
        try:
            data = resp.json()
        except ValueError:
            return "HTTP %s: %s" % (resp.status_code, (resp.text or '')[:300])
        if isinstance(data, dict) and data.get('description'):
            return "HTTP %s - %s (%s): %s" % (resp.status_code, data.get('errorName'),
                                              data.get('code'), data['description'])
        return "HTTP %s: %s" % (resp.status_code, str(data)[:300])

    def _get_access_token(self, force=False):
        """Return a (cached) app-only access token for the Partner Center API."""
        self.ensure_one()
        rec = self.sudo()
        now = fields.Datetime.now()
        if (not force and rec.access_token and rec.token_expiry
                and rec.token_expiry > now + timedelta(minutes=5)):
            return rec.access_token
        if not (rec.partner_tenant_id and rec.client_id and rec.client_secret):
            raise M365GraphError(_("Partner Directory ID, Client ID and Client Secret are "
                                   "required."))
        url = "%s/%s/oauth2/v2.0/token" % (rec.authority_url.rstrip('/'),
                                           rec.partner_tenant_id.strip())
        try:
            resp = requests.post(url, data={
                'client_id': rec.client_id.strip(),
                'client_secret': rec.client_secret.strip(),
                'scope': rec.api_scope.strip(),
                'grant_type': 'client_credentials',
            }, timeout=30)
        except requests.RequestException as exc:
            raise M365GraphError(_("Unable to reach the Microsoft identity platform: %s", exc))
        if resp.status_code != 200:
            raise M365GraphError(_("Authentication failed. %s",
                                   self.env['m365.tenant']._format_error(resp)))
        payload = resp.json()
        rec.write({
            'access_token': payload['access_token'],
            'token_expiry': now + timedelta(seconds=int(payload.get('expires_in', 3600))),
        })
        return rec.access_token

    def _pc_request(self, method, path, json_body=None, request_id=None, expect=()):
        """Call the Partner Center REST API.

        :param request_id: MS-RequestId; send the one of a call that got no answer again
            so that Microsoft does not execute the operation twice
        :param expect: error statuses returned instead of raised (e.g. 404)
        :return: (HTTP status, JSON body or ``{}``)
        :raise M365PartnerCenterTimeout: no answer, retry with the same request ID
        """
        self.ensure_one()
        url = "%s/%s" % (self.api_url.rstrip('/'), path.lstrip('/'))
        request_id = request_id or str(uuid.uuid4())
        retries, refreshed = 0, False
        while True:
            headers = {
                'Authorization': 'Bearer %s' % self._get_access_token(),
                'Accept': 'application/json',
                'MS-RequestId': request_id,
                'MS-CorrelationId': str(uuid.uuid4()),
                'MS-PartnerCenter-Application': PC_APPLICATION,
                'X-Locale': 'en-US',
            }
            try:
                resp = requests.request(method, url, headers=headers, json=json_body,
                                        timeout=120)
            except requests.RequestException as exc:
                raise M365PartnerCenterTimeout(
                    _("Partner Center did not answer: %s", exc)) from None
            if resp.status_code == 401 and not refreshed:
                refreshed = True
                self._get_access_token(force=True)
                continue
            if resp.status_code in (429, 503) and retries < 3:
                retries += 1
                try:
                    wait = int(resp.headers.get('Retry-After', 2))
                except ValueError:
                    wait = 2
                time.sleep(min(max(wait, 1), 30))
                continue
            if resp.status_code >= 400 and resp.status_code not in expect:
                raise M365GraphError(self._format_error(resp))
            return resp.status_code, (resp.json() if resp.content else {})

    # ------------------------------------------------------------------
    # Buttons
    # ------------------------------------------------------------------
    def action_test_connection(self):
        self.ensure_one()
        try:
            _status, profile = self._pc_request('GET', '/v1/profiles/legalbusiness')
        except M365GraphError as exc:
            self.write({'state': 'error', 'last_error': str(exc)})
            return self.env['m365.tenant']._notify(
                _("Partner Center connection failed: %s", exc), 'danger', reload=True)
        self.write({'state': 'connected', 'last_error': False,
                    'partner_name': profile.get('companyName')})
        return self.env['m365.tenant']._notify(
            _("Connected to Partner Center as %s.", profile.get('companyName') or self.name))

    def action_view_provisions(self):
        self.ensure_one()
        action = self.env['ir.actions.act_window']._for_xml_id(
            'm365_license_management.action_m365_tenant_provision')
        action['domain'] = [('partner_account_id', '=', self.id)]
        action['context'] = {'default_partner_account_id': self.id}
        return action


class M365TenantProvision(models.Model):
    """A customer tenant to create in Microsoft 365 (one row of ``Customers.csv``)."""
    _name = 'm365.tenant.provision'
    _description = 'Microsoft 365 Tenant to Create'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'id desc'

    # Customer data (CompanyProfile / BillingProfile of the Partner Center customer)
    name = fields.Char(string="Company Name", required=True, tracking=True)
    domain = fields.Char(
        string="Domain", required=True, tracking=True, copy=False,
        help="Initial domain of the new tenant, e.g. contoso.onmicrosoft.com. "
             "\"contoso\" is completed with .onmicrosoft.com.")
    partner_account_id = fields.Many2one(
        'm365.partner.account', string="Partner Center Account", required=True,
        tracking=True, default=lambda self: self._default_partner_account())
    company_id = fields.Many2one('res.company', string="Company",
                                 default=lambda self: self.env.company)
    first_name = fields.Char(string="Contact First Name")
    last_name = fields.Char(string="Contact Last Name")
    email = fields.Char(string="Contact Email",
                        help="Billing contact of the customer at Microsoft.")
    phone = fields.Char(string="Phone")
    address_line1 = fields.Char(string="Address Line 1")
    address_line2 = fields.Char(string="Address Line 2")
    city = fields.Char(string="City")
    address_state = fields.Char(string="State / Region",
                                help="State or region code, e.g. KA or WA.")
    postal_code = fields.Char(string="Postal Code")
    country_id = fields.Many2one('res.country', string="Country")
    address_confirmed = fields.Boolean(
        string="Address Confirmed",
        help="Create the tenant with the address as entered even when Microsoft cannot "
             "verify it.")
    culture = fields.Char(string="Culture", help="e.g. en-US or en-IN.")
    language = fields.Char(string="Language", help="Two letter code, e.g. en.")
    organization_registration_number = fields.Char(
        string="Organization Registration Number",
        help="Required by Microsoft in some countries, e.g. India, Brazil or Poland.")
    enable_gdap = fields.Boolean(string="Create GDAP Relationship", default=True)

    # Processing
    state = fields.Selection([
        ('draft', 'To Check'),
        ('ready', 'Ready'),
        ('queued', 'Creating'),
        ('done', 'Created'),
        ('failed', 'Failed'),
    ], string="Status", default='draft', readonly=True, copy=False, tracking=True,
        index=True)
    address_status = fields.Char(string="Address Verification", readonly=True, copy=False)
    error_message = fields.Text(string="Error", readonly=True, copy=False)
    queued_by = fields.Many2one('res.users', string="Requested By", readonly=True,
                                copy=False)
    pc_request_id = fields.Char(
        string="Pending Request ID", readonly=True, copy=False,
        help="MS-RequestId of a creation that got no answer, sent again on retry so that "
             "the tenant is not created twice.")

    # Result
    customer_id = fields.Char(string="Customer ID", readonly=True, copy=False)
    customer_tenant_id = fields.Char(string="Directory (Tenant) ID", readonly=True,
                                     copy=False, index=True)
    created_date = fields.Datetime(string="Created On", readonly=True, copy=False)
    admin_username = fields.Char(string="Admin User", readonly=True, copy=False,
                                 groups=ADMIN_GROUP)
    admin_password = fields.Char(string="Initial Admin Password", readonly=True,
                                 copy=False, groups=ADMIN_GROUP)
    tenant_record_id = fields.Many2one('m365.tenant', string="Odoo Tenant",
                                       compute='_compute_tenant_record_id')

    _domain_uniq = models.Constraint('unique(domain)',
                                     'A tenant with this domain is already in the list.')

    @api.depends('customer_tenant_id')
    def _compute_tenant_record_id(self):
        Tenant = self.env['m365.tenant'].with_context(active_test=False)
        tenants = Tenant.search([('tenant_id', 'in', [t for t in self.mapped(
            'customer_tenant_id') if t])])
        by_id = {tenant.tenant_id: tenant for tenant in tenants}
        for rec in self:
            rec.tenant_record_id = by_id.get(rec.customer_tenant_id, Tenant)

    @api.model
    def _default_partner_account(self):
        return self.env['m365.partner.account'].search([], limit=1)

    @staticmethod
    def _normalise_domain(value):
        """``Contoso`` -> ``contoso.onmicrosoft.com``"""
        value = (value or '').strip().lower()
        if value and '.' not in value:
            value += INITIAL_DOMAIN
        return value

    @api.model
    def _prepare_values(self, vals):
        vals = dict(vals)
        if vals.get('domain'):
            vals['domain'] = self._normalise_domain(vals['domain'])
        return vals

    @api.model_create_multi
    def create(self, vals_list):
        Account = self.env['m365.partner.account']
        vals_list = [self._prepare_values(vals) for vals in vals_list]
        for vals in vals_list:
            account = (Account.browse(vals['partner_account_id'])
                       if vals.get('partner_account_id') else self._default_partner_account())
            if account:
                vals.setdefault('culture', account.default_culture)
                vals.setdefault('language', account.default_language)
                vals.setdefault('enable_gdap', account.enable_gdap)
        return super().create(vals_list)

    def write(self, vals):
        vals = self._prepare_values(vals)
        if CUSTOMER_FIELDS & set(vals):
            if self.filtered(lambda r: r.state in ('queued', 'done')):
                raise UserError(_("A tenant being created or already created cannot be "
                                  "changed."))
            if 'state' not in vals:
                # Checked again with the new data
                self.filtered(lambda r: r.state in ('ready', 'failed')).write(
                    {'state': 'draft', 'pc_request_id': False})
        return super().write(vals)

    @api.ondelete(at_uninstall=False)
    def _unlink_except_created(self):
        if self.filtered(lambda r: r.state in ('queued', 'done')):
            raise UserError(_("A tenant being created or already created in Microsoft 365 "
                              "cannot be deleted."))

    @api.onchange('partner_account_id')
    def _onchange_partner_account_id(self):
        account = self.partner_account_id
        if account:
            self.culture = self.culture or account.default_culture
            self.language = self.language or account.default_language
            self.enable_gdap = account.enable_gdap

    @api.onchange('domain')
    def _onchange_domain(self):
        self.domain = self._normalise_domain(self.domain) or False

    def _check_admin(self):
        if not self.env.user.has_group(ADMIN_GROUP):
            raise AccessError(_("Only Microsoft 365 administrators can perform this action."))

    def _audit(self, status, message):
        self.env['m365.audit.log'].sudo().create({
            'action': 'provision_tenant',
            'status': status,
            'message': message,
            'user_id': (self.queued_by or self.env.user).id,
        })

    def _commit(self):
        """Keep the result of Microsoft calls, which cannot be undone."""
        if not config['test_enable']:
            self.env.cr.commit()

    # ------------------------------------------------------------------
    # Checks
    # ------------------------------------------------------------------
    def _check_values(self):
        """Errors in the data of the customer, before calling Microsoft."""
        self.ensure_one()
        errors = [_("%s is required.", self._fields[fname].string)
                  for fname in REQUIRED_FIELDS if not self[fname]]
        if self.domain and not DOMAIN_RE.match(self.domain):
            errors.append(_("Domain \"%s\" is not valid: use letters, digits and hyphens "
                            "followed by .onmicrosoft.com, e.g. contoso.onmicrosoft.com.",
                            self.domain))
        if self.email and not email_normalize(self.email):
            errors.append(_("\"%s\" is not a valid email address.", self.email))
        if self.culture and not CULTURE_RE.match(self.culture):
            errors.append(_("Culture \"%s\" is not valid, e.g. en-US.", self.culture))
        if self.language and not LANGUAGE_RE.match(self.language):
            errors.append(_("Language \"%s\" is not valid, e.g. en.", self.language))
        if (self.country_id.code in REGISTRATION_COUNTRIES
                and not self.organization_registration_number):
            errors.append(_("Microsoft requires the organization registration number for "
                            "%s.", self.country_id.name))
        return errors

    def _address_payload(self):
        address = {
            'AddressLine1': self.address_line1,
            'AddressLine2': self.address_line2,
            'City': self.city,
            'State': self.address_state,
            'PostalCode': self.postal_code,
            'Country': self.country_id.code,
        }
        return {key: value for key, value in address.items() if value}

    def _check_with_microsoft(self):
        """Domain availability and address verification by Partner Center."""
        self.ensure_one()
        account = self.partner_account_id
        errors = []
        # 200: the domain exists, 404: it is available
        status, _data = account._pc_request('HEAD', '/v1/domains/%s' % self.domain,
                                            expect=(404,))
        if status != 404:
            errors.append(_("Domain %s is already used in Microsoft 365.", self.domain))
        if self.address_confirmed:
            return errors
        _status, result = account._pc_request('POST', '/v1/validations/address',
                                              json_body=self._address_payload())
        self.address_status = result.get('status')
        if self.address_status not in VERIFIED_ADDRESS:
            message = _("Microsoft could not verify the address (%(status)s). %(details)s",
                        status=self.address_status or _("no status"),
                        details=result.get('validationMessage') or '')
            suggested = (result.get('suggestedAddresses') or [{}])[0]
            if suggested:
                message += _(" Suggested address: %s.", ", ".join(
                    suggested[key] for key in ('addressLine1', 'addressLine2', 'city',
                                               'state', 'postalCode', 'country')
                    if suggested.get(key)))
            errors.append(message + _(" Correct the address, or tick Address Confirmed to "
                                      "use it as entered."))
        return errors

    def action_check(self):
        """Check the data, the domain availability and the address."""
        self._check_admin()
        todo = self.filtered(lambda r: r.state in ('draft', 'ready'))
        for rec in todo:
            errors = rec._check_values()
            if not errors:
                try:
                    errors = rec._check_with_microsoft()
                except M365GraphError as exc:
                    errors = [str(exc)]
            rec.write({'state': 'draft' if errors else 'ready',
                       'error_message': "\n".join(errors) or False})
        ready = todo.filtered(lambda r: r.state == 'ready')
        message = _("%(ready)s of %(count)s tenant(s) ready to be created.",
                    ready=len(ready), count=len(todo))
        if ready != todo:
            message += " " + _("Open the others for the errors.")
        return self.env['m365.tenant']._notify(message,
                                               'success' if ready == todo else 'warning')

    # ------------------------------------------------------------------
    # Creation
    # ------------------------------------------------------------------
    def action_create(self):
        """Queue the creation; the scheduled action creates the tenants one by one."""
        self._check_admin()
        todo = self.filtered(lambda r: r.state in ('ready', 'failed'))
        if not todo:
            raise UserError(_("Check the tenants first: only checked or failed tenants can "
                              "be created."))
        todo.write({'state': 'queued', 'error_message': False, 'queued_by': self.env.uid})
        cron = self.env.ref('m365_license_management.ir_cron_m365_provision_tenants',
                            raise_if_not_found=False)
        if cron:
            cron.sudo()._trigger()
        return self.env['m365.tenant']._notify(
            _("%s tenant(s) queued. They are created in Microsoft 365 in the background, "
              "in about a minute.", len(todo)))

    @api.model
    def _cron_provision_tenants(self):
        queued = self.search([('state', '=', 'queued')], order='id',
                             limit=PROVISION_BATCH + 1)
        for rec in queued[:PROVISION_BATCH]:
            rec._provision()
        if len(queued) > PROVISION_BATCH:
            self.env.ref('m365_license_management.ir_cron_m365_provision_tenants')._trigger()

    def _customer_payload(self):
        address = dict(self._address_payload(), FirstName=self.first_name,
                       LastName=self.last_name)
        if self.phone:
            address['PhoneNumber'] = self.phone
        company = {'Domain': self.domain}
        if self.organization_registration_number:
            company['OrganizationRegistrationNumber'] = self.organization_registration_number
        return {
            'enableGDAPByDefault': self.enable_gdap,
            'CompanyProfile': company,
            'BillingProfile': {
                'Culture': self.culture,
                'Email': self.email,
                'Language': self.language,
                'CompanyName': self.name,
                'DefaultAddress': address,
            },
        }

    def _provision(self):
        """Create the customer tenant. Failures are stored, never raised."""
        self.ensure_one()
        if not self.pc_request_id:
            # Saved before the call: when the call gets no answer, the retry sends the same
            # ID and Microsoft returns the customer it created instead of a second one
            self.pc_request_id = str(uuid.uuid4())
            self._commit()
        try:
            _status, data = self.partner_account_id._pc_request(
                'POST', '/v1/customers', json_body=self._customer_payload(),
                request_id=self.pc_request_id)
        except M365GraphError as exc:
            vals = {'state': 'failed', 'error_message': str(exc)}
            if not isinstance(exc, M365PartnerCenterTimeout):
                vals['pc_request_id'] = False
            self.write(vals)
            self._audit('failed', "%s (%s): %s" % (self.name, self.domain, exc))
            self._commit()
            return False
        profile = data.get('companyProfile') or {}
        credentials = data.get('userCredentials') or {}
        domain = profile.get('domain') or self.domain
        username = credentials.get('userName') or ''
        if username and '@' not in username:
            username = '%s@%s' % (username, domain)
        self.write({
            'state': 'done',
            'error_message': False,
            'pc_request_id': False,
            'customer_id': data.get('id'),
            'customer_tenant_id': profile.get('tenantId') or data.get('id'),
            'created_date': fields.Datetime.now(),
            'admin_username': username or False,
            'admin_password': credentials.get('password') or False,
        })
        self.message_post(body=_("Tenant created in Microsoft 365. Directory (tenant) ID: "
                                 "%s.", self.customer_tenant_id))
        self._audit('success', _("%(name)s (%(domain)s) created, customer ID %(id)s.",
                                 name=self.name, domain=domain, id=self.customer_id))
        self._commit()
        return True

    # ------------------------------------------------------------------
    # Other buttons
    # ------------------------------------------------------------------
    def action_reset_draft(self):
        self.filtered(lambda r: r.state == 'failed').write(
            {'state': 'draft', 'pc_request_id': False})
        return True

    def action_clear_password(self):
        """Forget the initial admin password once it is handed over to the customer."""
        self._check_admin()
        for rec in self.filtered(lambda r: r.sudo().admin_password):
            rec.sudo().admin_password = False
            rec.message_post(body=_("Initial admin password removed from Odoo."))
        return True

    def action_register_tenant(self):
        """Open the Odoo tenant of the new customer, to manage its licenses."""
        self.ensure_one()
        if self.tenant_record_id:
            return {'type': 'ir.actions.act_window', 'res_model': 'm365.tenant',
                    'res_id': self.tenant_record_id.id, 'view_mode': 'form'}
        return {
            'type': 'ir.actions.act_window',
            'name': _("New Tenant"),
            'res_model': 'm365.tenant',
            'view_mode': 'form',
            'target': 'current',
            'context': {
                'default_name': self.name,
                'default_domain': self.domain,
                'default_tenant_id': self.customer_tenant_id,
                'default_company_id': self.company_id.id,
            },
        }


# Columns of Customers.csv: (field, header, required). Headers are matched ignoring the
# case, spaces and underscores, so the field labels can be used as well.
IMPORT_COLUMNS = [
    ('name', 'CompanyName', True),
    ('domain', 'Domain', True),
    ('first_name', 'FirstName', True),
    ('last_name', 'LastName', True),
    ('email', 'Email', True),
    ('phone', 'PhoneNumber', False),
    ('address_line1', 'AddressLine1', True),
    ('address_line2', 'AddressLine2', False),
    ('city', 'City', True),
    ('address_state', 'State', False),
    ('postal_code', 'PostalCode', True),
    ('country_id', 'Country', True),
    ('culture', 'Culture', False),
    ('language', 'Language', False),
    ('organization_registration_number', 'OrganizationRegistrationNumber', False),
]
IMPORT_TEMPLATE = 'm365_customers_template.xlsx'
IMPORT_SAMPLE = ('Contoso Ltd', 'contosoltd', 'Adele', 'Vance', 'adele.vance@contoso.com',
                 '+91 80 4000 1000', '1 MG Road', '', 'Bengaluru', 'KA', '560001', 'IN',
                 'en-IN', 'en', 'U72900KA2020PTC000000')


class M365TenantProvisionImportWizard(models.TransientModel):
    """Load the customers to create from ``Customers.csv`` (one tenant per row)."""
    _name = 'm365.tenant.provision.import.wizard'
    _inherit = 'm365.file.import.mixin'
    _description = 'Import Microsoft 365 Tenants to Create'

    partner_account_id = fields.Many2one(
        'm365.partner.account', string="Partner Center Account", required=True,
        default=lambda self: self.env['m365.tenant.provision']._default_partner_account())
    file = fields.Binary(string="File", attachment=False,
                         help="CSV or Excel (.xlsx) file with one customer per row.")
    filename = fields.Char()
    template_file = fields.Binary(attachment=False)
    line_ids = fields.One2many('m365.tenant.provision.import.line', 'wizard_id',
                               string="Rows")
    line_count = fields.Integer(compute='_compute_counts')
    error_count = fields.Integer(compute='_compute_counts')

    @api.depends('line_ids.status')
    def _compute_counts(self):
        for wiz in self:
            wiz.line_count = len(wiz.line_ids)
            wiz.error_count = len(wiz.line_ids.filtered(lambda l: l.status == 'error'))

    @api.onchange('file', 'partner_account_id')
    def _onchange_file(self):
        self.line_ids = [Command.clear()]

    @staticmethod
    def _column_key(title):
        return re.sub(r'[\s_]', '', str(title)).lower()

    def _map_header(self, header):
        Provision = self.env['m365.tenant.provision']
        lookup = {}
        for fname, title, _required in IMPORT_COLUMNS:
            for name in (fname, title, Provision._fields[fname].string):
                lookup[self._column_key(name)] = fname
        mapping, unknown = {}, []
        for index, title in enumerate(header):
            if not title:
                continue
            fname = lookup.get(self._column_key(title))
            if not fname:
                unknown.append(title)
            elif fname in mapping.values():
                raise UserError(_("The column \"%s\" appears more than once.", title))
            else:
                mapping[index] = fname
        if unknown:
            raise UserError(_("Unknown column(s): %(unknown)s.\nAllowed columns: %(allowed)s.",
                              unknown=", ".join(unknown),
                              allowed=", ".join(title for _f, title, _r in IMPORT_COLUMNS)))
        missing = [title for fname, title, required in IMPORT_COLUMNS
                   if required and fname not in mapping.values()]
        if missing:
            raise UserError(_("Missing required column(s): %s.", ", ".join(missing)))
        return mapping

    def _parse(self):
        """Read and check the uploaded file.

        :return: list of dicts with the serial number (1 = first row below the header),
                 the m365.tenant.provision values and the errors of each data row
        """
        self.ensure_one()
        rows = self._read_rows()
        if len(rows) < 2:
            raise UserError(_("The file contains no customer rows below the header."))
        mapping = self._map_header(rows[0])
        Provision = self.env['m365.tenant.provision']
        countries = {}
        for country in self.env['res.country'].search([]):
            countries[country.code.lower()] = countries[country.name.lower()] = country
        # The domain is unique among the tenants to create of all companies
        existing = set(Provision.sudo().search([]).mapped('domain'))
        seen, result = {}, []
        for sn, values in enumerate(rows[1:], start=1):
            raw = {fname: values[index] if index < len(values) else ''
                   for index, fname in mapping.items()}
            vals = {fname: text for fname, text in raw.items() if text}
            vals['partner_account_id'] = self.partner_account_id.id
            vals.setdefault('culture', self.partner_account_id.default_culture)
            vals.setdefault('language', self.partner_account_id.default_language)
            vals['enable_gdap'] = self.partner_account_id.enable_gdap
            country = countries.get(raw.get('country_id', '').lower())
            vals['country_id'] = country.id if country else False
            vals = Provision._prepare_values(vals)
            errors = Provision.new(vals)._check_values()
            if raw.get('country_id') and not country:
                required = _("%s is required.", Provision._fields['country_id'].string)
                errors = [_("Country \"%s\" not found: use the ISO code, e.g. IN.",
                            raw['country_id'])] + [e for e in errors if e != required]
            domain = vals.get('domain')
            if domain in existing:
                errors.append(_("Domain %s is already in the tenants to create.", domain))
            elif domain in seen:
                errors.append(_("Duplicate Domain, same as SN. %s.", seen[domain]))
            seen.setdefault(domain, sn)
            result.append({'sn': sn, 'raw': raw, 'vals': vals, 'errors': errors})
        return result

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'name': _("Import Tenants to Create"),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_validate(self):
        self.ensure_one()
        self.line_ids = [Command.clear()] + [Command.create({
            'sn': row['sn'],
            'name': row['raw'].get('name'),
            'domain': row['vals'].get('domain'),
            'email': row['raw'].get('email'),
            'country': row['raw'].get('country_id'),
            'status': 'error' if row['errors'] else 'valid',
            'message': "\n".join(row['errors']) or False,
        }) for row in self._parse()]
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        if not self.line_ids:
            raise UserError(_("Validate the file before adding the tenants."))
        # Parse again: the file or the tenants to create may have changed since
        parsed = self._parse()
        invalid = [row for row in parsed if row['errors']]
        if invalid:
            raise UserError(_(
                "No tenant has been added. Correct these rows and validate the file "
                "again:\n%s", "\n".join("SN. %s: %s" % (row['sn'], " ".join(row['errors']))
                                         for row in invalid)))
        records = self.env['m365.tenant.provision'].create([row['vals'] for row in parsed])
        action = self.env['ir.actions.act_window']._for_xml_id(
            'm365_license_management.action_m365_tenant_provision')
        action.update(name=_("Imported Tenants"), domain=[('id', 'in', records.ids)],
                      target='main')
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _("Microsoft 365"),
                'message': _("%s tenant(s) added. Select them and use Check, then Create "
                             "in Microsoft 365.", len(records)),
                'type': 'success',
                'next': action,
            },
        }

    def action_download_template(self):
        self.ensure_one()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = _("Customers")
        required_fill = PatternFill('solid', fgColor='FFF2CC')
        Provision = self.env['m365.tenant.provision']
        for col, (fname, title, required) in enumerate(IMPORT_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = Font(bold=True)
            if required:
                cell.fill = required_fill
            hint = Provision._fields[fname].help or Provision._fields[fname].string
            cell.comment = Comment(_("Required. %s", hint) if required else hint, 'Odoo')
            sheet.column_dimensions[cell.column_letter].width = max(14, len(title) + 4)
        sheet.append(list(IMPORT_SAMPLE))
        sheet.freeze_panes = 'A2'
        return self._download_workbook(book, IMPORT_TEMPLATE)


class M365TenantProvisionImportLine(models.TransientModel):
    _name = 'm365.tenant.provision.import.line'
    _description = 'Microsoft 365 Tenant to Create Import Row'
    _order = 'sn'

    wizard_id = fields.Many2one('m365.tenant.provision.import.wizard', required=True,
                                ondelete='cascade')
    sn = fields.Integer(string="SN.")
    name = fields.Char(string="Company Name")
    domain = fields.Char(string="Domain")
    email = fields.Char(string="Contact Email")
    country = fields.Char(string="Country")
    status = fields.Selection([('valid', 'Valid'), ('error', 'Error')], string="Status")
    message = fields.Text(string="Errors")
