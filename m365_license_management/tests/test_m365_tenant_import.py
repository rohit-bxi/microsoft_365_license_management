# -*- coding: utf-8 -*-
import base64
import io

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

GUID_1 = '11111111-2222-3333-4444-555555555555'
GUID_2 = '66666666-7777-8888-9999-000000000000'
GUID_3 = 'aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee'
GUID_4 = '12345678-1234-1234-1234-123456789012'
GUID_5 = 'abcdefab-cdef-abcd-efab-cdefabcdefab'
HEADER = ('Tenant Name,Primary Domain,Directory (Tenant) ID,Application (Client) ID,'
          'Client Secret,Enable License Operations,Low Availability Threshold (%)\n')


@tagged('post_install', '-at_install')
class TestM365TenantImport(TransactionCase):

    def _wizard(self, content, filename='tenants.csv'):
        if isinstance(content, str):
            content = content.encode()
        return self.env['m365.tenant.import.wizard'].create({
            'file': base64.b64encode(content), 'filename': filename,
        })

    def test_import_csv(self):
        wizard = self._wizard(HEADER
                              + 'Contoso,contoso.com,%s,%s,s1,no,15\n' % (GUID_1, GUID_2)
                              + '\n'
                              + 'Fabrikam,fabrikam.com,%s,%s,,,\n' % (GUID_2.upper(), GUID_3))
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (2, 0))
        # Serial numbers start at 1 below the header, blank rows are skipped
        self.assertEqual(wizard.line_ids.mapped('sn'), [1, 2])
        wizard.action_import()
        tenants = self.env['m365.tenant'].search([('tenant_id', 'in', (GUID_1, GUID_2))])
        self.assertEqual(len(tenants), 2)
        contoso = tenants.filtered(lambda t: t.name == 'Contoso')
        self.assertFalse(contoso.allow_license_operations)
        self.assertEqual(contoso.low_availability_percent, 15)
        self.assertEqual(contoso.sudo().client_secret, 's1')
        fabrikam = tenants - contoso
        # Blank or missing switches are TRUE, other blank cells keep the model defaults
        # and GUIDs are stored in lower case
        self.assertEqual(fabrikam.tenant_id, GUID_2)
        self.assertTrue(fabrikam.allow_license_operations)
        self.assertTrue(all(tenants.mapped('auto_approve_admins')))
        self.assertEqual(fabrikam.low_availability_percent, 10)
        self.assertEqual(fabrikam.graph_url, 'https://graph.microsoft.com/v1.0')
        self.assertFalse(wizard.file)

    def test_invalid_rows_create_nothing(self):
        Tenant = self.env['m365.tenant'].with_context(active_test=False)
        Tenant.create({'name': 'Existing', 'domain': 'existing.com', 'tenant_id': GUID_3,
                       'client_id': GUID_1, 'active': False})
        count = Tenant.search_count([])
        wizard = self._wizard(HEADER
                              + 'Good,good.com,%s,%s,,,\n' % (GUID_1, GUID_2)
                              + ',bad domain,not-a-guid,%s,,maybe,150\n' % GUID_2
                              + 'Dup,dup.com,%s,%s,,,\n' % (GUID_1, GUID_2)
                              + 'Old,old.com,%s,%s,,,\n' % (GUID_3.upper(), GUID_2)
                              + 'GOOD,Good.com,%s,%s,,,\n' % (GUID_4, GUID_2)
                              + 'existing,EXISTING.COM,%s,%s,,,\n' % (GUID_5, GUID_2))
        with self.assertRaisesRegex(UserError, 'Validate the file'):
            wizard.action_import()
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (6, 5))
        good, bad, dup, old, same, existing = wizard.line_ids
        self.assertEqual(good.status, 'valid')
        for text in ('Tenant Name is required', 'not a valid domain', 'not a valid ID',
                     'Yes/No', 'between 0 and 100'):
            self.assertIn(text, bad.message)
        # Duplicates within the file and with existing (even archived) tenants,
        # ignoring the case; the client ID may be shared by several tenants
        self.assertEqual(dup.message, 'Duplicate Directory (Tenant) ID, same as SN. 1.')
        self.assertEqual(old.message, 'Directory (Tenant) ID "%s" already exists in Odoo.'
                         % GUID_3)
        self.assertEqual(same.message, 'Duplicate Primary Domain, same as SN. 1.\n'
                                       'Duplicate Tenant Name, same as SN. 1.')
        self.assertEqual(existing.message, 'Primary Domain "existing.com" already exists '
                                           'in Odoo.\nTenant Name "existing" already exists '
                                           'in Odoo.')
        with self.assertRaisesRegex(UserError, '(?s)SN. 2: .*SN. 6: '):
            wizard.action_import()
        self.assertEqual(Tenant.search_count([]), count)

    def test_import_checks_again(self):
        wizard = self._wizard(HEADER + 'Contoso,contoso.com,%s,%s,,,\n' % (GUID_1, GUID_2))
        wizard.action_validate()
        self.assertEqual(wizard.error_count, 0)
        # Created by someone else after the validation
        self.env['m365.tenant'].create({'name': 'Contoso Group', 'domain': 'contoso.com',
                                        'tenant_id': GUID_3, 'client_id': GUID_2})
        with self.assertRaisesRegex(UserError, 'SN. 1: Primary Domain "contoso.com"'):
            wizard.action_import()
        self.assertFalse(self.env['m365.tenant'].search([('tenant_id', '=', GUID_1)]))

    def test_bad_header(self):
        wizard = self._wizard('Tenant Name,Primary Domain,Tenant\nA,a.com,x\n')
        with self.assertRaisesRegex(UserError, 'Unknown column'):
            wizard.action_validate()
        wizard = self._wizard('name,domain,tenant_id\nA,a.com,%s\n' % GUID_1)
        with self.assertRaisesRegex(UserError, 'Application \\(Client\\) ID'):
            wizard.action_validate()
        with self.assertRaisesRegex(UserError, 'Unsupported file type'):
            self._wizard('x', filename='tenants.txt').action_validate()

    def test_import_xlsx_from_template(self):
        wizard = self.env['m365.tenant.import.wizard'].create({})
        wizard.action_download_template()
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.template_file)))
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, [
            'Tenant Name', 'Primary Domain', 'Directory (Tenant) ID', 'Application (Client) ID',
            'Client Secret', 'Authority URL', 'Graph API URL', 'Graph Scope',
            'Enable License Operations', 'Skip Approval for Administrators',
            'Scheduled Synchronisation', 'Synchronise Users'])
        self.assertEqual(sheet.max_row, 5)  # header + 4 example tenants
        # Columns that are not in the template can still be added
        extra = ['Low Availability Threshold (Units)', 'Company']
        for col, title in enumerate(extra, start=len(headers) + 1):
            sheet.cell(row=1, column=col, value=title)
        row = dict.fromkeys(headers + extra)
        row.update({'Tenant Name': 'Woodgrove Bank', 'Primary Domain': 'woodgrovebank.com',
                    'Directory (Tenant) ID': GUID_1, 'Application (Client) ID': GUID_2,
                    'Scheduled Synchronisation': False,
                    'Low Availability Threshold (Units)': 5.0,
                    'Company': self.env.company.name})
        sheet.append(list(row.values()))
        out = io.BytesIO()
        book.save(out)
        wizard.write({'file': base64.b64encode(out.getvalue()), 'filename': 'tenants.xlsx'})
        wizard.action_validate()
        # The example tenants are valid test data
        self.assertEqual((wizard.line_count, wizard.error_count), (5, 0))
        wizard.action_import()
        Tenant = self.env['m365.tenant']
        fabrikam = Tenant.search([('domain', '=', 'fabrikam.com')])
        self.assertEqual((fabrikam.allow_license_operations, fabrikam.sync_users),
                         (False, False))
        northwind = Tenant.search([('domain', '=', 'northwindtraders.com')])
        self.assertEqual(northwind.graph_url, 'https://graph.microsoft.us/v1.0')
        adventure = Tenant.search([('domain', '=', 'adventure-works.com')])
        self.assertFalse(adventure.sudo().client_secret)
        self.assertTrue(adventure.allow_license_operations and adventure.auto_approve_admins)
        tenant = Tenant.search([('tenant_id', '=', GUID_1)])
        self.assertEqual(tenant.name, 'Woodgrove Bank')
        self.assertFalse(tenant.sync_enabled)
        self.assertTrue(tenant.sync_users and tenant.allow_license_operations)
        self.assertEqual(tenant.authority_url, 'https://login.microsoftonline.com')
        self.assertEqual(tenant.low_availability_units, 5)
        self.assertEqual(tenant.company_id, self.env.company)
